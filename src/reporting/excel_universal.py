from __future__ import annotations
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference

def write_management_workbook(tables:dict[str,pd.DataFrame],output):
    output=Path(output); output.parent.mkdir(parents=True,exist_ok=True)
    with pd.ExcelWriter(output,engine='openpyxl') as w:
        for name,df in tables.items(): df.to_excel(w,sheet_name=name[:31],index=False)
    wb=load_workbook(output)
    for ws in wb.worksheets:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True); cell.alignment=Alignment(horizontal='center')
        for col in ws.columns:
            letter=col[0].column_letter; width=min(40,max(12,max(len(str(x.value or '')) for x in col)+2)); ws.column_dimensions[letter].width=width
    wb.save(output); return output
