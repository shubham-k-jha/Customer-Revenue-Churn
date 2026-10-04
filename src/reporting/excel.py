from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference

def build_excel(scored_csv, output_xlsx):
    df=pd.read_csv(scored_csv)
    kpis=pd.DataFrame({'KPI':['Customers','Churn rate','Monthly revenue','Revenue at risk'],'Value':[len(df),df['churn_flag'].mean(),df['MonthlyCharges'].sum(),df['monthly_revenue_at_risk'].sum()]})
    by_contract=df.groupby('Contract',dropna=False).agg(customers=('customerID','count'),churn_rate=('churn_flag','mean'),monthly_revenue=('MonthlyCharges','sum'),revenue_at_risk=('monthly_revenue_at_risk','sum')).reset_index()
    risk=df.sort_values('churn_probability',ascending=False).head(500)
    with pd.ExcelWriter(output_xlsx,engine='openpyxl') as w:
        kpis.to_excel(w,index=False,sheet_name='Executive KPI'); by_contract.to_excel(w,index=False,sheet_name='Contract Analysis'); risk.to_excel(w,index=False,sheet_name='Risk Queue')
    wb=load_workbook(output_xlsx)
    for ws in wb.worksheets:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78'); cell.alignment=Alignment(horizontal='center')
        for col in ws.columns:
            letter=col[0].column_letter; ws.column_dimensions[letter].width=min(max(max(len(str(c.value or '')) for c in col)+2,10),35)
    if 'Contract Analysis' in wb.sheetnames:
        ws=wb['Contract Analysis']; chart=BarChart(); chart.title='Churn Rate by Contract'; chart.y_axis.title='Rate'; chart.x_axis.title='Contract'; chart.add_data(Reference(ws,min_col=3,min_row=1,max_row=ws.max_row),titles_from_data=True); chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=ws.max_row)); ws.add_chart(chart,'G2')
    wb.save(output_xlsx); return Path(output_xlsx)
